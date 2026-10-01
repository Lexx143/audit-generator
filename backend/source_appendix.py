"""Paginate original content into the corporate presentation, without an LLM.

Tables remain editable. Raster image bytes and aspect ratios are preserved.
Long paragraphs/rows are continued on as many slides as needed.
"""
import io

from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_AUTO_SIZE, MSO_ANCHOR
from pptx.util import Inches, Pt

import prototypes

FONT = "Source Sans Pro"
RED = RGBColor(0xE3, 0x4A, 0x4E)
BLACK = RGBColor(0x22, 0x22, 0x22)


def wrap(text, width, size):
    """Conservative line width for a portable Source Sans Pro layout.

    Explicit line breaks and all characters are retained, including identifiers
    without spaces. The extra allowance covers font substitution in Office.
    """
    capacity = max(1, int(width / Pt(size * .65)))
    lines = []
    for paragraph in text.split("\n"):
        if not paragraph:
            lines.append("")
        while paragraph:
            end = min(len(paragraph), capacity)
            if end < len(paragraph):
                space = paragraph.rfind(" ", 0, end + 1)
                if space > 0:
                    end = space + 1
            lines.append(paragraph[:end])
            paragraph = paragraph[end:]
    return lines or [""]


def style(frame, text, size=11, bold=False, color=BLACK):
    frame.clear()
    frame.word_wrap = True
    frame.auto_size = MSO_AUTO_SIZE.NONE
    frame.margin_left = frame.margin_right = 0
    frame.margin_top = frame.margin_bottom = 0
    for i, line in enumerate(text.split("\n")):
        p = frame.paragraphs[0] if i == 0 else frame.add_paragraph()
        p.text = line
        p.font.name, p.font.size, p.font.bold = FONT, Pt(size), bold
        p.font.color.rgb = color
        p.space_before = p.space_after = 0
        p.line_spacing = Pt(size * 1.35)


class Appendix:
    def __init__(self, prs, template, clone_slide):
        self.prs, self.template, self.clone_slide = prs, template, clone_slide
        self.left = Inches(.5)
        self.width = prs.slide_width - self.left * 2
        self.top, self.bottom = Inches(2.45), prs.slide_height - Inches(.65)
        self.slide = None
        self.y = self.top
        self.section = "Исходные материалы"
        self.filename = ""

    def box(self, text, left, top, width, height, size=11, bold=False, color=BLACK):
        shape = self.slide.shapes.add_textbox(left, int(top), int(width), int(height))
        style(shape.text_frame, text, size, bold, color)
        return shape

    def page(self):
        self.slide = self.clone_slide(self.prs, self.template)
        # Letterhead lives in the original slide layout/master.
        for shape in list(self.slide.shapes):
            shape._element.getparent().remove(shape._element)
        title_size = 19
        display_title = self.section if len(self.section) <= 150 else self.section[:147] + "…"
        title_lines = wrap(display_title, self.width, title_size)
        if len(title_lines) > 3:
            title_size = 14
            title_lines = wrap(display_title, self.width, title_size)
        title_height = Pt(title_size * 1.35 * len(title_lines))
        self.box(display_title, self.left, Inches(1.5), self.width, title_height, title_size, True)
        self.y = max(self.top, Inches(1.5) + title_height + Inches(.2))
        footer = f"{self.filename} · исходные материалы"
        self.box(footer, self.left, self.bottom + Inches(.12), self.width, Inches(.4), 8,
                 color=RGBColor(0x77, 0x77, 0x77))

    def paragraph(self, text, heading=False):
        size = 13 if heading else 11
        line_height = Pt(size * 1.35)
        lines = wrap(text, self.width, size)
        # Preserve text exactly when it fits; only inject breaks on continuation.
        if self.slide is None:
            self.page()
        height = len(lines) * line_height + Pt(10)
        if height <= self.bottom - self.y:
            self.box(text, self.left, self.y, self.width, height, size, heading)
            self.y += height
            return
        while lines:
            available = int((self.bottom - self.y - Pt(10)) / line_height)
            if available < 2:
                self.page()
                available = int((self.bottom - self.y - Pt(10)) / line_height)
            chunk, lines = lines[:available], lines[available:]
            height = len(chunk) * line_height + Pt(10)
            self.box("\n".join(chunk), self.left, self.y, self.width, height, size, heading)
            self.y += height
            if lines:
                self.page()

    def picture(self, document, block, new_page=True):
        if new_page or self.slide is None:
            self.page()
        if block.get("caption"):
            self.paragraph(block["caption"], heading=True)
        width, height = block["width"], block["height"]
        scale = min(self.width / width, (self.bottom - self.y) / height)
        w, h = int(width * scale), int(height * scale)
        self.slide.shapes.add_picture(io.BytesIO(prototypes.image_bytes(document, block)),
                                      self.left + (self.width - w) // 2, int(self.y), w, h)
        self.slide = None

    def table(self, rows):
        if not rows:
            return
        cols = max(len(row) for row in rows)
        rows = [row + [""] * (cols - len(row)) for row in rows]
        # Very wide tables become field/value records, without dropping columns.
        if cols > 8 or (cols >= 4 and any(len(cell) > 350 for row in rows for cell in row)):
            for number, row in enumerate(rows[1:], 1):
                self.paragraph(f"Строка {number}", heading=True)
                self.table([["Поле", "Значение"]] + [[label, value] for label, value in zip(rows[0], row)])
            return
        size = 10.5
        widths = [self.width // cols] * cols
        if cols == 2 and rows[0][0].strip().lower() in {"поле", "параметр"}:
            widths = [int(self.width * .28), int(self.width * .72)]
        elif cols > 2 and rows[0][0].strip().lower() in {"id", "№"}:
            first = Inches(.4)
            widths = [first] + [(self.width - first) // (cols - 1)] * (cols - 1)
        line_height = Pt(size * 1.35)
        padding = Pt(12)
        line_rows = [[wrap(cell, width - Pt(12), size) for cell, width in zip(row, widths)] for row in rows]
        header = rows[0]
        header_height = max(len(c) for c in line_rows[0]) * line_height + padding
        # Exceptionally verbose column headers are represented as field/value rows.
        if header_height > (self.bottom - self.top) / 3 and cols > 2:
            for row in rows[1:]:
                self.table([["Поле", "Значение"]] + [[label, value] for label, value in zip(header, row)])
            return
        if header_height > (self.bottom - self.top) / 3:
            # All text remains, even for a two-cell table with prose as its first row.
            for number, row in enumerate(rows, 1):
                self.paragraph(f"Строка {number}", heading=True)
                for cell in row:
                    self.paragraph(cell)
            return
        pending, heights = [header], [header_height]
        if self.slide is None or self.bottom - self.y < header_height + Pt(65):
            self.page()

        def flush():
            nonlocal pending, heights
            self.draw_table(pending, widths, heights, size)
            pending, heights = [header], [header_height]

        for row, remaining in zip(rows[1:], line_rows[1:]):
            first = True
            while any(remaining):
                available = int((self.bottom - self.y - sum(heights) - padding) / line_height)
                if available < 1:
                    flush()
                    self.page()
                    continue
                count = max(len(cell) for cell in remaining)
                if count <= available:
                    values = row if first else ["\n".join(cell) for cell in remaining]
                    pending.append(values)
                    heights.append(count * line_height + padding)
                    break
                # Avoid splitting an ordinary row if it fits on a fresh page.
                fresh = int((self.bottom - self.top - header_height - padding) / line_height)
                if count <= fresh and len(pending) > 1:
                    flush()
                    self.page()
                    continue
                pending.append(["\n".join(cell[:available]) for cell in remaining])
                heights.append(available * line_height + padding)
                remaining = [cell[available:] for cell in remaining]
                first = False
                flush()
                self.page()
        if len(pending) > 1 or len(rows) == 1:
            flush()
        self.y += Pt(14)

    def draw_table(self, rows, widths, heights, size):
        table = self.slide.shapes.add_table(len(rows), len(widths), self.left, int(self.y),
                                             self.width, int(sum(heights))).table
        for column, width in zip(table.columns, widths):
            column.width = int(width)
        for i, row in enumerate(rows):
            table.rows[i].height = int(heights[i])
            for j, value in enumerate(row):
                cell = table.cell(i, j)
                cell.margin_left = cell.margin_right = Pt(6)
                cell.margin_top = cell.margin_bottom = Pt(5)
                cell.vertical_anchor = MSO_ANCHOR.TOP
                cell.fill.solid()
                cell.fill.fore_color.rgb = RED if i == 0 else RGBColor.from_string("FFF7F7" if i % 2 == 0 else "FFFFFF")
                style(cell.text_frame, value, size, i == 0, RGBColor(255, 255, 255) if i == 0 else BLACK)
        self.y += sum(heights)


def append_sources(prs, documents, clone_slide):
    if not documents:
        return
    # At this point the final two slides are conclusions and corporate contacts.
    contacts = prs.slides._sldIdLst[-1]
    appendix = Appendix(prs, prs.slides[-2], clone_slide)
    for document in documents:
        appendix.filename = document["filename"]
        appendix.section = "Материалы обследования"
        appendix.page()
        appendix.paragraph(document["filename"], heading=True)
        appendix.paragraph("Полный перенос исходных текстов, таблиц и изображений. "
                           "Исходные формулировки и незаполненные поля сохранены; "
                           "оценка и рекомендации приведены в первой части отчета.")
        pending_headings = []
        for block in document["blocks"]:
            kind = block["kind"]
            if kind == "heading":
                pending_headings.append(block["text"])
                continue
            new_section = bool(pending_headings)
            if pending_headings:
                appendix.section = pending_headings[-1]
                appendix.page()
                for heading in pending_headings[:-1]:
                    appendix.paragraph(heading, heading=True)
                if len(pending_headings[-1]) > 150:
                    appendix.paragraph(pending_headings[-1], heading=True)
                pending_headings = []
            if kind == "paragraph":
                appendix.paragraph(block["text"])
            elif kind == "table":
                appendix.table(block["rows"])
            elif kind == "image":
                appendix.picture(document, block, new_page=not new_section)
        for heading in pending_headings:
            appendix.paragraph(heading, heading=True)
    prs.slides._sldIdLst.remove(contacts)
    prs.slides._sldIdLst.append(contacts)
