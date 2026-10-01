import asyncio
import base64
import hashlib
import io
import json
from pathlib import Path
from unittest.mock import AsyncMock, Mock
from zipfile import ZipFile

import pytest
from docx import Document
from fastapi import FastAPI
from fastapi.testclient import TestClient
from PIL import Image
from pptx import Presentation
from pptx.util import Inches

import llm
import pptx_builder
import prototypes
from schemas import AuditData, AuditStructure, Case, ParseRequest, ReviseRequest
from source_appendix import append_sources
from routers import audit, images, pptx, prototypes as upload


def picture():
    stream = io.BytesIO()
    Image.new('RGB', (120, 40), 'navy').save(stream, 'PNG')
    return stream.getvalue()


def source():
    doc = Document()
    doc.add_heading('Сеть филиала', level=1)
    doc.add_paragraph('Исходное описание, которое нельзя потерять.')
    doc.add_paragraph('☐ Резервирование проверено')
    table = doc.add_table(rows=1, cols=3)
    for cell, text in zip(table.rows[0].cells, ['Узел', 'IP', 'Состояние']):
        cell.text = text
    for i in range(48):
        for cell, text in zip(table.add_row().cells, [f'Узел-{i}', f'10.0.0.{i}', 'Down' if i == 0 else 'Up']):
            cell.text = text
    doc.add_paragraph('Схема с оригинальной подписью')
    doc.add_picture(io.BytesIO(picture()))
    doc.add_picture(io.BytesIO(picture()))  # repeated placements must survive
    stream = io.BytesIO()
    doc.save(stream)
    return stream.getvalue()


@pytest.fixture(autouse=True)
def store(tmp_path, monkeypatch):
    monkeypatch.setattr(prototypes, 'STORE', tmp_path / 'prototypes')


@pytest.fixture
def client():
    app = FastAPI()
    for router in (upload.router, audit.router, pptx.router, images.router):
        app.include_router(router)
    return TestClient(app)


def test_extract_preserves_rows_checkboxes_and_repeated_images():
    document = prototypes.ingest('проект.docx', source())
    assert document['tables'] == 1 and document['images'] == 2
    assert len(next(b['rows'] for b in document['blocks'] if b['kind'] == 'table')) == 49
    assert '☐ Резервирование проверено' in prototypes.context([document['id']])
    assert prototypes.load(document['id']) == document
    for block in document['blocks']:
        if block['kind'] == 'image':
            assert prototypes.image_bytes(document, block) == picture()
    assert not any('base64' in str(block) for block in document['blocks'])


def test_rejects_invalid_and_oversized_documents(monkeypatch):
    for name, raw in [('a.pdf', b'%PDF'), ('a.docx', b'not a zip')]:
        with pytest.raises(prototypes.PrototypeError):
            prototypes.ingest(name, raw)
    monkeypatch.setattr(prototypes, 'MAX_FILE_BYTES', 10)
    with pytest.raises(prototypes.PrototypeError, match='20 МБ'):
        prototypes.ingest('a.docx', b'x' * 11)
    with pytest.raises(prototypes.PrototypeError):
        prototypes.load('../../secret')


def test_context_never_silently_truncates(monkeypatch):
    document = prototypes.ingest('a.docx', source())
    monkeypatch.setattr(prototypes, 'MAX_CONTEXT_CHARS', 40)
    with pytest.raises(prototypes.PrototypeError, match='не были обрезаны'):
        prototypes.context([document['id']])


def test_upload_atomic_and_unknown_source(client):
    result = client.post('/api/prototypes', files=[('files', ('a.docx', source())), ('files', ('b.pdf', b'no'))])
    assert result.status_code == 422
    assert not list(prototypes.STORE.glob('*/manifest.json'))
    good = client.post('/api/prototypes', files={'files': ('a.docx', source())})
    assert good.status_code == 200
    metadata = good.json()['documents'][0]
    assert 'blocks' not in metadata and metadata['images'] == 2
    assert client.get('/api/prototypes/' + metadata['id']).json() == metadata
    assert client.get('/api/prototypes/' + 'a' * 32).status_code == 404


def case(**changes):
    data = dict(title='Нет резерва', vulnerability='Один uplink', risk='Простой', recommendation='Резерв',
                priority='ПЕРВЫЙ ПРИОРИТЕТ', category='Сеть', image_prompt='')
    return Case(**(data | changes))


def test_parse_and_revision_keep_sources_outside_llm(monkeypatch):
    document = prototypes.ingest('a.docx', source())
    original = AuditStructure(client_name='Тест', review='Обзор', cases=[case()], conclusions=['Проверить'])
    mock = AsyncMock(return_value=original)
    monkeypatch.setattr(llm, '_parse_structured', mock)
    request = ParseRequest(general_data='', vulnerabilities='', conclusions='', audit_type='full',
                           source_ids=[document['id']], generate_illustrations=False)
    result = asyncio.run(llm.generate_structure(request))
    assert result.source_ids == [document['id']]
    prompt = mock.call_args.args[0]
    assert 'Узел-47' in prompt and 'не являются кейсами' in prompt
    assert 'Генерация иллюстраций отключена' in prompt
    assert mock.call_args.kwargs['response_format'] is AuditStructure
    assert 'недоверенные' in mock.call_args.kwargs['system']
    result.cases[0].image_b64 = 'data:image/png;base64,original'
    result.cases[0].image_source = 'uploaded'
    revised = asyncio.run(llm.revise_structure(ReviseRequest(current_data=result, revision_prompt='Уточни риск', audit_type='full')))
    assert revised.source_ids == result.source_ids
    assert revised.cases[0].image_source == 'uploaded'
    assert revised.cases[0].image_b64 == result.cases[0].image_b64
    assert document['id'] not in llm._strip_images(result)


def test_empty_image_prompts_do_not_swap_uploaded_images():
    a, b = case(title='A', image_b64='A', image_source='uploaded'), case(title='B', image_b64='B', image_source='uploaded')
    original = AuditData(client_name='Тест', review='', cases=[a, b], conclusions=[])
    result = llm._reattach_images(original.model_copy(deep=True), original)
    assert [c.image_b64 for c in result.cases] == ['A', 'B']


def test_disabled_generation_never_calls_provider(client, monkeypatch):
    build = AsyncMock()
    monkeypatch.setattr(llm, 'build_image_prompt', build)
    result = client.post('/api/generate_image', json={'prompt': 'test', 'vulnerability': 'test', 'generate_illustrations': False})
    assert result.status_code == 409
    build.assert_not_called()


def test_export_removes_only_automatic_images(client, monkeypatch, tmp_path):
    template = tmp_path / 'template.pptx'
    template.touch()
    monkeypatch.setattr(pptx_builder, 'TEMPLATE_PATH', str(template))
    build = Mock(return_value=io.BytesIO(b'fake export'))
    monkeypatch.setattr(pptx_builder, 'build_pptx', build)
    data = AuditData(client_name='Тест', review='', conclusions=[], source_ids=['b' * 32], cases=[
        case(image_b64='generated', image_source='generated'),
        case(image_b64='library', image_source='library'),
        case(image_b64='manual', image_source='uploaded', image_reusable=True),
        case(image_b64='legacy', image_reusable=True),
    ])
    result = client.post('/api/generate_pptx', json={'data': data.model_dump(), 'generate_illustrations': False})
    assert result.status_code == 200
    exported = build.call_args.args[0]
    assert [c.image_b64 for c in exported.cases] == [None, None, 'manual', None]
    assert exported.source_ids == data.source_ids
    assert data.cases[0].image_b64 == 'generated'


def test_appendix_preserves_native_tables_images_and_continuations():
    document = prototypes.ingest('a.docx', source())
    prs = Presentation()
    prs.slide_width, prs.slide_height = 7559675, 10691813
    prs.slides.add_slide(prs.slide_layouts[6])
    end = prs.slides.add_slide(prs.slide_layouts[6])
    end.shapes.add_textbox(0, 0, Inches(1), Inches(1)).text = 'Контакты'
    append_sources(prs, [document], pptx_builder.clone_slide)
    text = '\n'.join(s.text if s.has_text_frame else '\n'.join(c.text for row in s.table.rows for c in row.cells) if s.has_table else ''
                     for slide in prs.slides for s in slide.shapes)
    for i in range(48):
        assert f'Узел-{i}' in text and f'10.0.0.{i}' in text
    assert '☐ Резервирование проверено' in text
    assert prs.slides[-1].shapes[0].text == 'Контакты'
    pictures = [s for slide in prs.slides for s in slide.shapes if s.shape_type == 13]
    assert len(pictures) == 2
    assert all(s.image.blob == picture() for s in pictures)
    assert all(abs(s.width / s.height - 3) < .001 for s in pictures)
    assert sum(s.has_table for slide in prs.slides for s in slide.shapes) > 1
    assert all(s.top + s.height < prs.slide_height for slide in list(prs.slides)[1:-1] for s in slide.shapes)
    saved = io.BytesIO()
    prs.save(saved)
    assert len(Presentation(saved).slides) == len(prs.slides)


def test_very_long_rows_and_headings_terminate_without_losing_tail():
    document = prototypes.ingest('a.docx', source())
    document['blocks'] = [
        {'kind': 'heading', 'text': 'Длинный заголовок ' * 30 + 'КОНЕЦЗАГОЛОВКА'},
        {'kind': 'table', 'rows': [['Поле', 'Значение'], ['Настройки', 'Данные параметра ' * 700 + 'КОНЕЦТАБЛИЦЫ']]},
        {'kind': 'paragraph', 'text': 'Описание ' * 700 + 'КОНЕЦТЕКСТА'},
    ]
    prs = Presentation()
    prs.slide_width, prs.slide_height = 7559675, 10691813
    for _ in range(2):
        prs.slides.add_slide(prs.slide_layouts[6])
    append_sources(prs, [document], pptx_builder.clone_slide)
    text = '\n'.join(t for slide in prs.slides for t in slide._element.xpath('.//a:t/text()'))
    for tail in ['КОНЕЦЗАГОЛОВКА', 'КОНЕЦТАБЛИЦЫ', 'КОНЕЦТЕКСТА']:
        assert tail in text
    assert 3 < len(prs.slides) < 100


def test_word_automatic_numbering_is_kept():
    doc = Document()
    doc.add_paragraph('Первый пункт', style='List Number')
    doc.add_paragraph('Второй пункт', style='List Number')
    doc.add_paragraph('Маркер', style='List Bullet')
    buffer = io.BytesIO()
    doc.save(buffer)
    parsed = prototypes.ingest('lists.docx', buffer.getvalue())
    texts = [block['text'] for block in parsed['blocks']]
    assert texts == ['1. Первый пункт', '2. Второй пункт', '• Маркер']


def test_appending_after_deleted_template_slides_has_unique_zip_parts(tmp_path, monkeypatch):
    prs = Presentation()
    prs.slide_width, prs.slide_height = 7559675, 10691813
    for _ in range(18):
        prs.slides.add_slide(prs.slide_layouts[6])
    template = tmp_path / 'template.pptx'
    prs.save(template)
    monkeypatch.setattr(pptx_builder, 'TEMPLATE_PATH', str(template))
    document = prototypes.ingest('a.docx', source())
    data = AuditData(client_name='Test', review='', cases=[], conclusions=[], source_ids=[document['id']])
    output = pptx_builder.build_pptx(data)
    with ZipFile(output) as archive:
        names = archive.namelist()
        assert len(names) == len(set(names))
    assert len(Presentation(output).slides) > 7


def test_header_text_and_images_are_not_lost():
    doc = Document()
    doc.add_paragraph('Основной текст')
    header = doc.sections[0].header.paragraphs[0]
    header.text = 'Название объекта в колонтитуле'
    header.add_run().add_picture(io.BytesIO(picture()))
    buffer = io.BytesIO()
    doc.save(buffer)
    parsed = prototypes.ingest('header.docx', buffer.getvalue())
    assert parsed['images'] == 1
    assert 'Название объекта в колонтитуле' in prototypes.context([parsed['id']])
    assert any(block.get('caption') == 'Название объекта в колонтитуле' for block in parsed['blocks'])
